"""Builds the import fixtures for folio-io's spreadsheet and presentation tests.

Run with a Python that has openpyxl and python-pptx:
    python make_fixtures.py
Then, when LibreOffice is at hand, `make_fixtures.sh` converts them so the tests also read
files written by LibreOffice (with saved formula results, shared formulas, ODS and ODP).
"""
import datetime
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))


def budget():
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Household budget 2026"
    ws["A1"].font = Font(bold=True, size=16)
    ws.merge_cells("A1:E1")
    head = ["Item", "Category", "Amount", "Date", "Share"]
    for i, h in enumerate(head):
        c = ws.cell(row=2, column=i + 1, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=Side(style="thin"))
    rows = [
        ("Rent", "Home", 1200, datetime.date(2026, 1, 1)),
        ("Groceries", "Food", 450.5, datetime.date(2026, 1, 3)),
        ("Electricity", "Home", 85.25, datetime.date(2026, 1, 10)),
        ("Cinema", "Fun", 32, datetime.date(2026, 1, 14)),
        ("Train pass", "Travel", 79.9, datetime.date(2026, 1, 20)),
    ]
    for r, (item, cat, amount, date) in enumerate(rows, start=3):
        ws.cell(row=r, column=1, value=item)
        ws.cell(row=r, column=2, value=cat)
        ws.cell(row=r, column=3, value=amount).number_format = "#,##0.00"
        d = ws.cell(row=r, column=4, value=date)
        d.number_format = "yyyy-mm-dd"
        s = ws.cell(row=r, column=5, value=f"=C{r}/$C$9")
        s.number_format = "0.0%"
    ws["A9"] = "Total"
    ws["A9"].font = Font(bold=True, italic=True)
    ws["C9"] = "=SUM(C3:C7)"
    ws["C9"].number_format = "#,##0.00"
    ws["C9"].font = Font(bold=True)
    ws["A10"] = "Biggest"
    ws["C10"] = "=MAX(C3:C7)"
    ws["A11"] = "Over budget?"
    ws["C11"] = '=IF(C9>Limits!B2,"yes","no")'
    ws["A12"] = "Home rate"
    ws["C12"] = "=VLOOKUP(\"Home\",'Rates & Notes'!A2:B5,2,FALSE)"
    ws["C12"].number_format = "0%"
    ws["A13"] = "Days"
    ws["C13"] = "=D7-D3"
    ws["A14"] = "Code"
    ws["C14"] = "00123"
    ws["A15"] = "Flag"
    ws["C15"] = True
    ws["A16"] = "Owner"
    ws["C16"] = "=\"Bob's \"&'Bob''s list'!A1"
    ws["A17"] = "Mean"
    ws["C17"] = "=AVERAGE(C3:C7)"
    ws["C17"].font = Font(color="C00000", underline="single", strike=True)
    ws["C17"].alignment = Alignment(horizontal="right", wrap_text=True)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["D"].width = 14
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = "A2:E7"
    chart = BarChart()
    chart.type = "col"
    chart.title = "Spending"
    data = Reference(ws, min_col=3, min_row=2, max_row=7)
    cats = Reference(ws, min_col=1, min_row=3, max_row=7)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "G3")

    rates = wb.create_sheet("Rates & Notes")
    rates.append(["Category", "Rate"])
    for cat, rate in [("Home", 0.4), ("Food", 0.25), ("Fun", 0.1), ("Travel", 0.25)]:
        rates.append([cat, rate])
    rates["D1"] = "Notes"
    rates["D2"] = "Rates are shares of the monthly income."
    rates["B7"] = "=SUM(B2:B5)"

    limits = wb.create_sheet("Limits")
    limits["A1"] = "Limit"
    limits["A2"] = "Monthly"
    limits["B2"] = 2000

    bob = wb.create_sheet("Bob's list")
    bob["A1"] = "Ann"

    trend = wb.create_sheet("Trend")
    trend.append(["Month", "Sales", "Costs"])
    for i, m in enumerate(["Jan", "Feb", "Mar", "Apr"]):
        trend.append([m, 10 + i * 3, 6 + i])
    line = LineChart()
    line.title = "Sales and costs"
    line.add_data(Reference(trend, min_col=2, max_col=3, min_row=1, max_row=5), titles_from_data=True)
    line.set_categories(Reference(trend, min_col=1, min_row=2, max_row=5))
    trend.add_chart(line, "E2")
    pie = PieChart()
    pie.add_data(Reference(trend, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    pie.set_categories(Reference(trend, min_col=1, min_row=2, max_row=5))
    trend.add_chart(pie, "E20")

    wb.defined_names["Income"] = DefinedName("Income", attr_text="Limits!$B$2")
    wb.save(os.path.join(HERE, "budget.xlsx"))


def deck():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    # 1: title slide (placeholders inherit their positions from the layout).
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Quarterly review"
    s.placeholders[1].text = "October 2026"
    s.notes_slide.notes_text_frame.text = "Welcome everyone."
    # 2: title and content with bullets and runs.
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Highlights"
    body = s.placeholders[1].text_frame
    body.text = "Revenue up 12%"
    p = body.add_paragraph()
    p.text = "New office"
    p.level = 1
    p = body.add_paragraph()
    r = p.add_run()
    r.text = "Bold "
    r.font.bold = True
    r = p.add_run()
    r.text = "and red"
    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r.font.italic = True
    r.font.size = Pt(28)
    # 3: blank with shapes, a picture, a table and a chart.
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5), Inches(3), Inches(1.5))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    box.line.color.rgb = RGBColor(0, 0, 0)
    box.line.width = Pt(2)
    box.text_frame.text = "Boxed"
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    oval = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(0.5), Inches(2), Inches(2))
    oval.rotation = 30
    s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.5), Inches(0.5), Inches(2), Inches(1))
    tb = s.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(4), Inches(1))
    tb.text_frame.text = "A text box in Courier"
    tb.text_frame.paragraphs[0].runs[0].font.name = "Courier New"
    png = os.path.join(HERE, "dot.png")
    s.shapes.add_picture(png, Inches(9), Inches(0.5), Inches(1), Inches(1))
    rows, cols = 3, 3
    t = s.shapes.add_table(rows, cols, Inches(0.5), Inches(4), Inches(6), Inches(1.5)).table
    for r in range(rows):
        for c in range(cols):
            t.cell(r, c).text = f"r{r}c{c}"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("Sales", (10, 12, 15))
    cd.add_series("Costs", (7, 8, 9))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(3), Inches(5), Inches(3.5), cd)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(0xF4, 0xF1, 0xEA)
    # 4: hidden slide.
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Backup"
    s._element.set("show", "0")
    prs.save(os.path.join(HERE, "deck.pptx"))


def dot():
    # A 4x4 PNG so the deck has a picture.
    import struct
    import zlib

    w = h = 4
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")
    open(os.path.join(HERE, "dot.png"), "wb").write(png)


if __name__ == "__main__":
    dot()
    budget()
    deck()
